"""Servicio de Procesamiento y Normalización de Respaldos Avanzados (BackupProcessor).

Parsea los archivos TXT y artefactos generados por BackupService (contactos, SMS, llamadas,
calendario, aplicaciones, Wi-Fi, Bluetooth, ajustes/sistema, diccionario y portapapeles).
Genera:
- Una estructura normalizada en memoria y en JSON (`Procesado/backup_procesado.json`).
- Un libro Excel (`Procesado/backup_procesado.xlsx`) con pestañas separadas por categoría:
  Contactos, SMS, Llamadas, Calendario, Aplicaciones, Wi-Fi, Bluetooth,
  Ajustes y sistema, Diccionario, Portapapeles, Resumen.
- Mantiene siempre intactos los archivos TXT originales.
- Si una categoría o archivo es irregular (dumpsys, errores, permisos denegados),
  extrae campos reconocibles y preserva el texto original.
"""

from __future__ import annotations

import json
import logging
import os
import re
import sys
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple
from xml.sax.saxutils import escape as xml_escape

logger = logging.getLogger("ui.backup_processor")


@dataclass
class CategoryProcessedData:
    id: str
    name: str
    status: str  # OK | PARTIAL | EMPTY | DENIED | ERROR | UNKNOWN
    total_records: int = 0
    empty_records: int = 0
    error_count: int = 0
    denied_count: int = 0
    raw_files: List[str] = field(default_factory=list)
    records: List[Dict[str, Any]] = field(default_factory=list)
    raw_fallback: Optional[str] = None
    note: Optional[str] = None

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)


@dataclass
class BackupProcessedResult:
    backup_dir: str
    processed_at: str
    overall_status: str
    categories: Dict[str, CategoryProcessedData] = field(default_factory=dict)
    json_path: Optional[str] = None
    xlsx_path: Optional[str] = None

    def to_dict(self) -> Dict[str, Any]:
        return {
            "backup_dir": self.backup_dir,
            "processed_at": self.processed_at,
            "overall_status": self.overall_status,
            "json_path": self.json_path,
            "xlsx_path": self.xlsx_path,
            "categories": {cat_id: cat_data.to_dict() for cat_id, cat_data in self.categories.items()},
        }


# =============================================================================
# Parsers especializados por categoría
# =============================================================================

def parse_content_query_rows(text: str) -> List[Dict[str, str]]:
    """Parsea salidas típicas de `content query --uri ...`.
    Formato general:
    Row: 0 _id=1, display_name=Juan, data1=+123456, ...
    o
    Row 0: name=John Doe, phone=+123456
    """
    rows: List[Dict[str, str]] = []
    if not text or not text.strip():
        return rows

    # Regex para capturar el inicio de cada fila
    row_chunks = re.split(r"(?:^|\n)\s*Row(?::|\s+)\d+[\s:]*", text.strip())
    for chunk in row_chunks:
        chunk = chunk.strip()
        if not chunk:
            continue
        row_dict: Dict[str, str] = {}
        # Parsear clave=valor separados por coma o espacio
        # Maneja comas dentro o separadas: key=val, key2=val2
        matches = re.findall(r'([a-zA-Z0-9_\-\.]+)=((?:(?!\s+[a-zA-Z0-9_\-\.]+=)[^,\n])*)', chunk)
        if matches:
            for k, v in matches:
                row_dict[k.strip()] = v.strip().rstrip(",")
        else:
            # Fallback simple
            parts = [p.strip() for p in chunk.split(",") if "=" in p]
            for p in parts:
                k, sep, v = p.partition("=")
                if sep:
                    row_dict[k.strip()] = v.strip()

        if row_dict:
            rows.append(row_dict)
        elif chunk and not chunk.startswith("["):
            rows.append({"raw": chunk})
    return rows


def parse_contacts(files_map: Dict[str, str]) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    # Prioridad: contacts_raw.txt o contacts_all.txt
    for fname in ("contacts_raw.txt", "contacts_all.txt", "contacts.txt"):
        if fname in files_map:
            text = files_map[fname]
            rows = parse_content_query_rows(text)
            for r in rows:
                name = r.get("display_name") or r.get("name") or r.get("display_name_alt") or ""
                phone = r.get("data1") or r.get("phone") or r.get("number") or ""
                rec_id = r.get("_id") or r.get("contact_id") or ""
                if name or phone:
                    records.append({
                        "ID": rec_id,
                        "Nombre": name,
                        "Teléfono": phone,
                        "Tipo": r.get("data2") or "",
                        "Detalle": json.dumps(r, ensure_ascii=False) if len(r) > 3 else "",
                    })
            if records:
                break
    return records


def parse_sms(files_map: Dict[str, str]) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    for fname in ("sms_all.txt", "sms_inbox.txt", "sms_sent.txt", "sms.txt"):
        if fname in files_map:
            text = files_map[fname]
            rows = parse_content_query_rows(text)
            for r in rows:
                addr = r.get("address") or r.get("phone") or ""
                body = r.get("body") or r.get("text") or ""
                date_val = r.get("date") or ""
                date_str = ""
                if date_val.isdigit():
                    try:
                        ms = int(date_val)
                        if ms > 1e11:
                            ms = ms / 1000.0
                        date_str = datetime.fromtimestamp(ms).strftime("%Y-%m-%d %H:%M:%S")
                    except Exception:
                        date_str = date_val
                else:
                    date_str = date_val

                tipo_sms = "Enviado" if r.get("type") == "2" else ("Recibido" if r.get("type") == "1" else (fname.replace("sms_", "").replace(".txt", "")))
                if addr or body:
                    records.append({
                        "ID": r.get("_id") or "",
                        "Remitente/Destino": addr,
                        "Mensaje": body,
                        "Fecha": date_str,
                        "Tipo": tipo_sms,
                    })
    return records


def parse_calls(files_map: Dict[str, str]) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    for fname in ("call_log.txt", "calls.txt"):
        if fname in files_map:
            text = files_map[fname]
            rows = parse_content_query_rows(text)
            for r in rows:
                num = r.get("number") or ""
                name = r.get("name") or r.get("cached_name") or ""
                date_val = r.get("date") or ""
                date_str = ""
                if date_val.isdigit():
                    try:
                        ms = int(date_val)
                        if ms > 1e11:
                            ms = ms / 1000.0
                        date_str = datetime.fromtimestamp(ms).strftime("%Y-%m-%d %H:%M:%S")
                    except Exception:
                        date_str = date_val
                duration = r.get("duration") or "0"
                call_type = r.get("type") or ""
                type_desc = {"1": "Entrante", "2": "Saliente", "3": "Perdida", "5": "Rechazada"}.get(call_type, call_type)
                records.append({
                    "Número": num,
                    "Nombre": name,
                    "Fecha": date_str,
                    "Duración (s)": duration,
                    "Tipo": type_desc,
                })
    return records


def parse_calendar(files_map: Dict[str, str]) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    for fname in ("calendar_events.txt", "calendar_list.txt", "cal.txt"):
        if fname in files_map:
            text = files_map[fname]
            rows = parse_content_query_rows(text)
            for r in rows:
                title = r.get("title") or r.get("calendar_displayName") or r.get("name") or ""
                desc = r.get("description") or ""
                dtstart = r.get("dtstart") or ""
                date_str = ""
                if dtstart.isdigit():
                    try:
                        ms = int(dtstart)
                        if ms > 1e11:
                            ms = ms / 1000.0
                        date_str = datetime.fromtimestamp(ms).strftime("%Y-%m-%d %H:%M:%S")
                    except Exception:
                        date_str = dtstart
                loc = r.get("eventLocation") or ""
                records.append({
                    "Título": title,
                    "Descripción": desc,
                    "Fecha Inicio": date_str,
                    "Ubicación": loc,
                    "Origen": fname,
                })
    return records


def parse_apps(files_map: Dict[str, str]) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    # Prioridad: apps_user.txt, apps_all.txt, apps_system.txt
    user_pkgs = set()
    if "apps_user.txt" in files_map:
        for line in files_map["apps_user.txt"].splitlines():
            line = line.strip().replace("package:", "")
            if line:
                user_pkgs.add(line)

    system_pkgs = set()
    if "apps_system.txt" in files_map:
        for line in files_map["apps_system.txt"].splitlines():
            line = line.strip().replace("package:", "")
            if line:
                system_pkgs.add(line)

    all_pkgs = set()
    for fname in ("apps_all.txt", "apps.txt"):
        if fname in files_map:
            for line in files_map[fname].splitlines():
                line = line.strip().replace("package:", "")
                if line:
                    all_pkgs.add(line)

    combined = user_pkgs | system_pkgs | all_pkgs
    for pkg in sorted(combined):
        if pkg in user_pkgs:
            tipo = "Usuario (3rd-party)"
        elif pkg in system_pkgs:
            tipo = "Sistema"
        else:
            tipo = "Instalada"
        records.append({
            "Paquete": pkg,
            "Tipo": tipo,
        })
    return records


def parse_wifi(files_map: Dict[str, str]) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    # 1. wifi_networks.txt (cmd wifi list-saved-networks)
    # Formato habitual: Network Id   SSID   Security type
    # o lista de redes
    if "wifi_networks.txt" in files_map:
        lines = files_map["wifi_networks.txt"].splitlines()
        for line in lines:
            line = line.strip()
            if not line or "Network Id" in line or line.startswith("===") or line.startswith("["):
                continue
            parts = re.split(r"\s{2,}|\t+", line)
            if len(parts) >= 2:
                records.append({
                    "ID": parts[0],
                    "SSID": parts[1].strip('"'),
                    "Seguridad": parts[2] if len(parts) > 2 else "Desconocida",
                    "Detalle": "",
                })
            else:
                m = re.search(r'SSID:\s*"([^"]+)"', line) or re.search(r'SSID:\s*([^\s,]+)', line)
                if m:
                    records.append({
                        "ID": "-",
                        "SSID": m.group(1),
                        "Seguridad": "Configurada",
                        "Detalle": line,
                    })

    # 2. wifi_dump.txt fallback si no hay redes parseadas
    if not records and "wifi_dump.txt" in files_map:
        dump_text = files_map["wifi_dump.txt"]
        ssids = set(re.findall(r'SSID:\s*"([^"]+)"', dump_text))
        for ssid in sorted(ssids):
            records.append({
                "ID": "-",
                "SSID": ssid,
                "Seguridad": "Detectada en dump",
                "Detalle": "",
            })
    return records


def parse_bluetooth(files_map: Dict[str, str]) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    text = files_map.get("bluetooth_devices.txt") or files_map.get("bluetooth.txt") or ""
    if text:
        # Busca bloques como Name: ..., Address: ..., Bond State: ... o dispositivos vinculados
        devices = re.findall(r'([0-9A-Fa-f]{2}(?::[0-9A-Fa-f]{2}){5})\s+\[(.*?)\]', text)
        for mac, name in devices:
            records.append({
                "Nombre": name.strip(),
                "MAC / Dirección": mac.strip(),
                "Estado": "Vinculado",
            })
        if not records:
            # Líneas con MAC
            mac_lines = re.findall(r'([0-9A-Fa-f]{2}(?::[0-9A-Fa-f]{2}){5})', text)
            for mac in sorted(set(mac_lines)):
                records.append({
                    "Nombre": "Dispositivo Bluetooth",
                    "MAC / Dirección": mac,
                    "Estado": "Detectado",
                })
    return records


def parse_settings(files_map: Dict[str, str]) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    for scope in ("global", "system", "secure"):
        fname = f"settings_{scope}.txt"
        if fname in files_map:
            for line in files_map[fname].splitlines():
                line = line.strip()
                if not line or line.startswith("["):
                    continue
                k, sep, v = line.partition("=")
                if sep:
                    records.append({
                        "Ámbito": scope.capitalize(),
                        "Clave": k.strip(),
                        "Valor": v.strip(),
                    })
    # Batería y display como info adicional
    if "battery_config.txt" in files_map:
        for line in files_map["battery_config.txt"].splitlines():
            line = line.strip()
            if ":" in line and not line.startswith("["):
                k, sep, v = line.partition(":")
                records.append({
                    "Ámbito": "Batería",
                    "Clave": k.strip(),
                    "Valor": v.strip(),
                })
    return records


def parse_dictionary(files_map: Dict[str, str]) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    for fname in ("user_dictionary.txt", "words.txt"):
        if fname in files_map:
            text = files_map[fname]
            rows = parse_content_query_rows(text)
            for r in rows:
                word = r.get("word") or ""
                freq = r.get("frequency") or ""
                locale = r.get("locale") or ""
                if word:
                    records.append({
                        "Palabra": word,
                        "Frecuencia": freq,
                        "Idioma/Locale": locale,
                    })
    return records


def parse_clipboard(files_map: Dict[str, str]) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    if "clipboard_current.txt" in files_map:
        curr = files_map["clipboard_current.txt"].strip()
        if curr and not curr.startswith("["):
            records.append({
                "Origen": "Actual",
                "Contenido": curr,
            })
    if "clipboard_history.txt" in files_map:
        hist = files_map["clipboard_history.txt"].strip()
        if hist and not hist.startswith("["):
            # Extraer textos citados o clips
            clips = re.findall(r'text\s*=\s*"(.*?)"', hist, re.DOTALL)
            for clip in clips:
                records.append({
                    "Origen": "Historial",
                    "Contenido": clip.strip(),
                })
    return records


# =============================================================================
# Clase principal: BackupProcessor
# =============================================================================

class BackupProcessor:
    """Procesador integral de respaldos Android."""

    CATEGORIES = [
        ("contacts", "Contactos", "Contacts", parse_contacts),
        ("sms", "SMS", "SMS", parse_sms),
        ("calls", "Llamadas", "Calls", parse_calls),
        ("calendar", "Calendario", "Calendar", parse_calendar),
        ("apps", "Aplicaciones", "Apps", parse_apps),
        ("wifi", "Wi-Fi", "WiFi", parse_wifi),
        ("bluetooth", "Bluetooth", "Bluetooth", parse_bluetooth),
        ("settings", "Ajustes y Sistema", "SystemSettings", parse_settings),
        ("dictionary", "Diccionario de Usuario", "Dictionary", parse_dictionary),
        ("clipboard", "Portapapeles", "Clipboard", parse_clipboard),
    ]

    def __init__(self, backup_dir: str | Path) -> None:
        self.backup_dir = Path(backup_dir)
        self.processed_dir = self.backup_dir / "Procesado"

    def process(self) -> BackupProcessedResult:
        """Procesa todas las carpetas del backup, genera JSON y XLSX, y retorna el resultado."""
        self.processed_dir.mkdir(parents=True, exist_ok=True)

        manifest = self._load_manifest()
        overall_status = manifest.get("OverallStatus", "UNKNOWN") if manifest else "UNKNOWN"
        processed_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        cat_results: Dict[str, CategoryProcessedData] = {}

        # 1. Procesar categorías estructuradas
        for cat_id, cat_name, sub_folder, parser_fn in self.CATEGORIES:
            cat_path = self.backup_dir / sub_folder
            files_map, raw_files = self._read_text_files(cat_path)

            status = "EMPTY"
            denied_count = 0
            error_count = 0
            empty_count = 0

            # Verificar si en los archivos hubo permisos denegados o errores
            for content in files_map.values():
                c_lower = content.lower()
                if "[permiso denegado]" in c_lower or "securityexception" in c_lower or "permission denial" in c_lower:
                    denied_count += 1
                elif "[error al ejecutar]" in c_lower:
                    error_count += 1
                elif not content.strip():
                    empty_count += 1

            records: List[Dict[str, Any]] = []
            fallback_text: Optional[str] = None
            note: Optional[str] = None

            if files_map:
                try:
                    records = parser_fn(files_map)
                except Exception as exc:
                    logger.warning("Error parseando categoría %s: %s", cat_id, exc)
                    error_count += 1

            # Si no se obtuvieron registros pero hay archivos, conservar texto completo como fallback
            if not records and files_map:
                all_text = []
                for fname, txt in files_map.items():
                    all_text.append(f"--- {fname} ---\n{txt}")
                fallback_text = "\n\n".join(all_text)
                note = "Esta categoría no pudo ser normalizada completamente en tabla; se conserva el texto original."

            # Determinar estado de la categoría
            if records:
                status = "OK" if denied_count == 0 and error_count == 0 else "PARTIAL"
            elif denied_count > 0:
                status = "DENIED"
            elif error_count > 0:
                status = "ERROR"
            elif empty_count > 0 or not files_map:
                status = "EMPTY"
            else:
                status = "OK"

            cat_results[cat_id] = CategoryProcessedData(
                id=cat_id,
                name=cat_name,
                status=status,
                total_records=len(records),
                empty_records=empty_count,
                error_count=error_count,
                denied_count=denied_count,
                raw_files=raw_files,
                records=records,
                raw_fallback=fallback_text,
                note=note,
            )

        # 2. Guardar JSON procesado
        json_path = self.processed_dir / "backup_procesado.json"
        res = BackupProcessedResult(
            backup_dir=str(self.backup_dir.resolve()),
            processed_at=processed_at,
            overall_status=overall_status,
            categories=cat_results,
            json_path=str(json_path.resolve()),
        )

        try:
            json_path.write_text(json.dumps(res.to_dict(), indent=2, ensure_ascii=False), encoding="utf-8")
        except Exception as exc:
            logger.error("No se pudo guardar backup_procesado.json: %s", exc)

        # 3. Guardar XLSX procesado
        xlsx_path = self.processed_dir / "backup_procesado.xlsx"
        try:
            self._generate_excel(xlsx_path, res)
            res.xlsx_path = str(xlsx_path.resolve())
        except Exception as exc:
            logger.error("No se pudo generar backup_procesado.xlsx: %s", exc)

        return res

    def _load_manifest(self) -> Optional[Dict[str, Any]]:
        for fname in ("MANIFEST.json", "RESUMEN_BACKUP.json"):
            path = self.backup_dir / fname
            if path.exists():
                try:
                    return json.loads(path.read_text(encoding="utf-8"))
                except Exception:
                    pass
        return None

    def _read_text_files(self, folder: Path) -> Tuple[Dict[str, str], List[str]]:
        """Lee todos los archivos .txt de la carpeta sin modificarlos."""
        files_map: Dict[str, str] = {}
        raw_files: List[str] = []
        if not folder.exists() or not folder.is_dir():
            return files_map, raw_files

        for file_path in folder.glob("*.txt"):
            rel_name = file_path.name
            raw_files.append(rel_name)
            try:
                files_map[rel_name] = file_path.read_text(encoding="utf-8", errors="replace")
            except Exception as exc:
                logger.warning("No se pudo leer archivo %s: %s", file_path, exc)
        return files_map, raw_files

    def _generate_excel(self, output_path: Path, result: BackupProcessedResult) -> None:
        """Crea el libro de Excel con las hojas especificadas."""
        sheets: List[Tuple[str, List[List[object]]]] = []

        # Hoja Resumen
        resumen_rows: List[List[object]] = [
            ["Categoría", "Estado", "Registros", "Archivos Originales", "Permisos Denegados", "Errores"],
        ]
        for cat_id, cat_data in result.categories.items():
            resumen_rows.append([
                cat_data.name,
                cat_data.status,
                cat_data.total_records,
                len(cat_data.raw_files),
                cat_data.denied_count,
                cat_data.error_count,
            ])
        sheets.append(("Resumen", resumen_rows))

        # Hojas por cada categoría
        for cat_id, cat_data in result.categories.items():
            sheet_title = cat_data.name[:31]  # Excel limita nombres de hoja a 31 chars
            rows: List[List[object]] = []
            if cat_data.records:
                # Encabezados
                headers = list(cat_data.records[0].keys())
                rows.append(headers)
                for rec in cat_data.records:
                    rows.append([rec.get(h, "") for h in headers])
            else:
                # Contenido de fallback o mensaje de estado
                rows.append(["Mensaje / Contenido Original"])
                if cat_data.raw_fallback:
                    for line in cat_data.raw_fallback.splitlines():
                        rows.append([line])
                else:
                    rows.append([f"Sin registros disponibles (Estado: {cat_data.status})"])
            sheets.append((sheet_title, rows))

        # Intentar con openpyxl primero
        try:
            from openpyxl import Workbook
            wb = Workbook()
            wb.remove(wb.active)  # remover hoja por defecto
            for name, rows in sheets:
                ws = wb.create_sheet(title=name)
                for r in rows:
                    ws.append(r)
                if ws.max_row > 1:
                    ws.freeze_panes = "A2"
                # Ajustar ancho de columnas
                for col in ws.columns:
                    col_letter = col[0].column_letter
                    max_len = max(len(str(c.value or "")) for c in col)
                    ws.column_dimensions[col_letter].width = min(60, max(12, max_len + 3))
            wb.save(output_path)
        except Exception:
            _write_minimal_xlsx(output_path, sheets)


def _excel_col_name(n: int) -> str:
    res = ""
    while n:
        n, rem = divmod(n - 1, 26)
        res = chr(65 + rem) + res
    return res


def _write_minimal_xlsx(path: Path, sheets: Sequence[Tuple[str, Sequence[Sequence[object]]]]) -> None:
    """Escritor mínimo de XLSX mediante zipfile si openpyxl no está disponible."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            + "".join(f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' for i in range(1, len(sheets) + 1))
            + "</Types>",
        )
        archive.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            '</Relationships>',
        )
        sheet_tags = "".join(f'<sheet name="{xml_escape(name)}" sheetId="{i}" r:id="rId{i}"/>' for i, (name, _) in enumerate(sheets, 1))
        archive.writestr(
            "xl/workbook.xml",
            f'<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>{sheet_tags}</sheets></workbook>',
        )
        rels = "".join(f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>' for i in range(1, len(sheets) + 1))
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rels}</Relationships>',
        )
        for idx, (_, rows) in enumerate(sheets, 1):
            xml_rows = []
            for r_idx, r in enumerate(rows, 1):
                cells = "".join(
                    f'<c r="{_excel_col_name(c_idx)}{r_idx}" t="inlineStr"><is><t>{xml_escape(str(val if val is not None else ""))}</t></is></c>'
                    for c_idx, val in enumerate(r, 1)
                )
                xml_rows.append(f'<row r="{r_idx}">{cells}</row>')
            archive.writestr(
                f"xl/worksheets/sheet{idx}.xml",
                '<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
                + "".join(xml_rows)
                + "</sheetData></worksheet>",
            )
